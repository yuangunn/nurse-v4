# 102병동 배정표 빈 양식 생성 — 병동이 준 .xls(데이터 가득)에서 구조만 떠서 만든다.
#   python3 scripts/make-form-102.py "standalone/102병동 어싸인 (최종).xls"
# .xls 원본은 실제 간호사 이름이 들어 있어 리포에 커밋하지 않는다.
import sys, xlrd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

src = sys.argv[1] if len(sys.argv) > 1 else 'standalone/102병동 어싸인 (최종).xls'
out = sys.argv[2] if len(sys.argv) > 2 else 'standalone/병실 배정표_102.xlsx'

sh = xlrd.open_workbook(src, formatting_info=True).sheet_by_index(0)
ROWS, COLS = 29, 17
NAME_R = range(3, 16)      # 4~16행 = 이름 칸 (0-based 3~15)
DAY_C = range(3, 10)       # D~J = 일~토

wb = Workbook(); ws = wb.active; ws.title = '배정표'
thin = Side(style='thin', color='BFBFBF')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
head = PatternFill('solid', fgColor='EEF3FE')

for c, ci in sh.colinfo_map.items():
    if c < COLS: ws.column_dimensions[get_column_letter(c+1)].width = round(ci.width/256, 2)
for r, ri in sh.rowinfo_map.items():
    if r < ROWS: ws.row_dimensions[r+1].height = round(ri.height/20, 1)

for r in range(ROWS):
    for c in range(COLS):
        v = sh.cell_value(r, c)
        if isinstance(v, float): v = int(v) if v == int(v) else v
        v = str(v).strip()
        # 이름·교육행사 내용은 비운다 — 그 주에만 맞는 값이 양식에 박히면 안 된다 (122 교훈)
        if v and (r in NAME_R or r == 16) and c in DAY_C: v = ''
        cell = ws.cell(row=r+1, column=c+1, value=v or None)
        cell.border = box
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.font = Font(name='맑은 고딕', size=9)

for (r1, r2, c1, c2) in sh.merged_cells:
    if r1 < ROWS and c1 < COLS:
        ws.merge_cells(start_row=r1+1, start_column=c1+1,
                       end_row=min(r2, ROWS), end_column=min(c2, COLS))

ws['A1'].font = Font(name='맑은 고딕', size=14, bold=True)
for c in range(1, 11): ws.cell(row=2, column=c).fill = head          # 요일 줄
for c in range(1, 11): ws.cell(row=2, column=c).font = Font(name='맑은 고딕', size=10, bold=True)
for r in (4, 10, 14): ws.cell(row=r, column=1).font = Font(name='맑은 고딕', size=12, bold=True)
ws.print_area = f'A1:{get_column_letter(COLS)}{ROWS}'
ws.page_setup.orientation = 'landscape'
ws.page_setup.fitToWidth = 1
wb.save(out)
print(f'{out} — {ROWS}행 × {COLS}열, 병합 {len(sh.merged_cells)}개, 이름 칸 비움')
